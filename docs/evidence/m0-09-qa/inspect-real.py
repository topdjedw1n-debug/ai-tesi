import asyncio,hashlib,json,re,zipfile
from pathlib import Path
from docx import Document as WordDocument
from sqlalchemy import select
from app.core.database import AsyncSessionLocal,get_engine
from app.models.document import DocumentSection
from app.services.ai_pipeline.text_utils import contains_concrete_evidence
async def main():
 q=Path('/qa');rows=json.loads((q/'real-results.json').read_text())+json.loads((q/'real-final-results.json').read_text());out=[]
 async with AsyncSessionLocal() as db:
  for row in rows:
   if row['status']!='completed':continue
   f=q/row['file'];assert hashlib.sha256(f.read_bytes()).hexdigest()==row['sha']
   with zipfile.ZipFile(f) as z:assert z.testzip() is None
   text='\n'.join(p.text for p in WordDocument(f).paragraphs)
   assert not re.findall(r'\[[A-Za-z][A-Za-z]+\d{4}[a-z]?\]',text)
   sections=(await db.execute(select(DocumentSection).where(DocumentSection.document_id==row['id'],DocumentSection.status=='completed'))).scalars().all()
   checks=[{'section':s.section_index,'checked':(s.claim_verification or {}).get('checked',0),'counts':(s.claim_verification or {}).get('counts',{}),'has_numeric_detail':contains_concrete_evidence(s.content or '')} for s in sections]
   assert all(x['checked']>0 or x['has_numeric_detail'] for x in checks),checks
   out.append({'id':row['id'],'file':row['file'],'sha256':row['sha'],'zip_valid':True,'unresolved_markers':0,'paragraphs':len(WordDocument(f).paragraphs),'characters':len(text),'final_semantic_grounding_guard_passes':True,'section_evidence':checks})
 await get_engine().dispose();(q/'real-docx-validation.json').write_text(json.dumps(out,indent=2));print(json.dumps(out))
asyncio.run(main())
