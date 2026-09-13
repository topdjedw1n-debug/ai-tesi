"""Export one live-written section from an S4 lab run as a small DOCX.

Usage: python scripts/lab_section_docx.py LAB_OUTPUT_DIR SECTION_INDEX OUT.docx [SECTION_EVIDENCE.json]
Reads the lab's replay.db (executor_section event of that section), resolves
[KEY] markers to in-text citations with the same formatter as S5, using the
pack recorded in the lab DB plus the lab's added sources, and writes a DOCX
with the section heading and paragraphs only (no bibliography).
"""

import json
import re
import sqlite3
import sys
from pathlib import Path

from docx import Document

MARKER = re.compile(r"\[(STD:[^\[\]\n]+|[\w:./-]+)\](?::\d+)?")
# The writer sometimes puts a fragment number outside the bracket ([KEY]:3);
# such a marker is resolved by its key family and the stray number dropped.


def main():
    out_dir, index, target = Path(sys.argv[1]), int(sys.argv[2]), Path(sys.argv[3])
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from app.services.ai_pipeline.citation_formatter import (
        CitationFormatter,
        CitationStyle,
    )

    db = sqlite3.connect(f"file:{out_dir / 'replay.db'}?mode=ro", uri=True)
    rows = db.execute(
        "select event_type, payload from document_provenance order by id"
    ).fetchall()
    events = [(et, json.loads(p)) for et, p in rows]
    section = next(
        p
        for et, p in events
        if et == "executor_section" and p.get("section_index") == index
    )
    pack = next(p for et, p in events if et == "executor_source_pack")
    inputs = next(p for et, p in events if et == "generation_replay_inputs")
    style = CitationStyle(inputs["executor_inputs"]["brief"]["citation_style"])
    known = {}
    for s in pack["sources"]:
        src = s["source"]
        known[s["citation_key"]] = (
            src.get("authors") or [src.get("title")],
            src.get("year"),
        )
    added = next(
        (
            p
            for et, p in events
            if et == "generation_provider_attempt"
            and p.get("section_index") == index
            and p.get("outcome") == "started"
        ),
        None,
    )
    if added:
        prompt = added["request"]["messages"][0]["content"]
        body = json.loads(prompt[prompt.index('{"forbidden_placeholders"') :])
        for e in body["evidence"]:
            known.setdefault(e["key"], None)
    lab_meta = {}
    if len(sys.argv) > 4:  # the section-evidence JSON used for the run
        for spec in json.load(open(sys.argv[4])).values():
            for item in spec.get("add") if isinstance(spec, dict) else spec:
                lab_meta[item["key"]] = (
                    item["authors"],
                    item["year"],
                    item.get("intext"),
                )
    text = section["raw_content"] if section.get("raw_content") else section["content"]
    unresolved = []

    def cite(m):
        key = m.group(1)
        meta = known.get(key) or lab_meta.get(key)
        if not meta:  # key family: KEY:1, KEY:2 ... share authors and year
            family = [k for k in lab_meta if k.split(":")[0] == key.split(":")[0]]
            meta = lab_meta.get(family[0]) if family else None
        if not meta:
            unresolved.append(key)
            return ""
        if len(meta) == 3 and meta[2]:
            return f"({meta[2]})"
        authors, year = meta[0], meta[1]
        return CitationFormatter.format_intext(authors, year, style=style)

    resolved = MARKER.sub(cite, text)
    resolved = re.sub(r"[ ]{2,}", " ", resolved)
    doc = Document()
    doc.add_heading(section["title"], level=1)
    words = 0
    for para in [p.strip() for p in resolved.split("\n") if p.strip()]:
        if para.startswith("## "):
            doc.add_heading(para[3:], level=2)
        else:
            doc.add_paragraph(para)
            words += len(para.split())
    doc.save(str(target))
    print(
        f"section {index} '{section['title']}' -> {target} ({words} words; unresolved markers: {sorted(set(unresolved))})"
    )


if __name__ == "__main__":
    main()
