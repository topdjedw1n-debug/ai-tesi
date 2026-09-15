"""Small Markdown-to-Word adapter for the markup emitted by the writer."""

from __future__ import annotations

import re
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

from app.services.ai_pipeline.citation_formatter import bibliography_heading

_PARSER = MarkdownIt("commonmark", {"html": False})
DEFAULT_DOCX_PROFILE = {
    "version": "academic-docx-v1",
    "basis": "system_default",
    "page_size": "A4",
    "font": "Times New Roman",
    "font_size_pt": 12,
    "line_spacing": 1.5,
    "margins_cm": 2.5,
}


def _heading_identity(value: str) -> str:
    value = re.sub(r"^\s*#{1,6}\s+", "", value).strip()
    value = re.sub(r"\s+#+\s*$", "", value)
    return " ".join(value.strip("*_ ").split()).replace("’", "'").casefold()


def assemble_section(title: str, content: str | None, level: int = 1) -> str:
    """Remove only repeated title lines at the beginning; preserve body headings."""
    lines = str(content or "").strip().splitlines()
    while lines and _heading_identity(lines[0]) == _heading_identity(title):
        lines.pop(0)
        while lines and not lines[0].strip():
            lines.pop(0)
    body = "\n".join(lines)
    return f"{'#' * max(1, min(int(level), 3))} {title}\n\n{body}".rstrip()


TOC_HEADINGS = {"it": "Indice", "en": "Contents", "uk": "Зміст"}
LEGAL_HEADINGS = {
    "it": "Normativa e giurisprudenza",
    "en": "Legislation and case law",
    "uk": "Нормативні акти та судова практика",
}
_FRAME_TITLES = re.compile(
    r"^\s*(?:introduzione|introduction|premessa|conclusioni|conclusion[s]?|"
    r"вступ|висновки|abstract|sommario|ringraziamenti)\b",
    re.I,
)
_NUMBERED = re.compile(
    r"^\s*(?:(?:capitolo|chapter|розділ)\s+\S+|\d+(?:\.\d+)*[.)]?\s)", re.I
)


def normalize_typography(text: str) -> str:
    """Word-like Italian typography: ’ inside words, «…» quotes, en dashes.

    A reviewer reads straight apostrophes and long dashes as machine output;
    the writer's markers and URLs are not touched (no spaces, no letters
    around the apostrophe in a DOI).
    """
    text = re.sub(r"(?<=\w)'\s?(?=\w)", "’", text)
    text = re.sub(r'"([^"\n]{1,400})"', r"«\1»", text)
    text = re.sub(r"\s*—\s*", " – ", text)
    return re.sub(r"[ \t]{2,}", " ", text)


def with_chapter_headings(
    sections: list[dict[str, Any]], nodes: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Insert a chapter heading (level 1, no body) before the first section of
    every root scope that no section represents itself: the plan maps
    sub-sections to the chapter's children and the reviewer still wants the
    chapter line and the 1.1 numbering under it."""
    child_ids = {c.get("scope_id") for n in nodes for c in n.get("children", [])}
    roots = [n for n in nodes if n.get("scope_id") not in child_ids]

    def descendants(node):
        return {node.get("scope_id")} | {
            d for c in node.get("children", []) for d in descendants(c)
        }

    out: list[dict[str, Any]] = []
    placed: set[str] = set()
    for section in sections:
        ids = set(section.get("scope_ids") or [])
        for root in roots:
            rid = root.get("scope_id")
            if rid in placed or not ids & descendants(root):
                continue
            placed.add(rid)
            if rid not in ids or int(section.get("level") or 1) != 1:
                out.append({"title": root["title"], "level": 1, "content": ""})
        out.append(section)
    return out


def assemble_document(
    sections: list[dict[str, Any]], bibliography: list[dict[str, Any]], language: str
) -> str:
    """The whole work as the Markdown the exporter renders: numbered chapters
    and sub-sections by plan level, normalized typography, an alphabetical
    bibliography and a separate list of legislation and case law."""
    parts: list[str] = []
    chapter = sub = 0
    for section in sections:
        level = max(1, min(int(section.get("level") or 1), 3))
        original = " ".join(str(section["title"]).split())
        title = original
        if level == 1 and not _FRAME_TITLES.match(title):
            # A supervisor-numbered chapter ("Capitolo I", "2.") keeps its title
            # but still counts, so its sub-sections are numbered under it.
            chapter, sub = chapter + 1, 0
            if not _NUMBERED.match(title):
                title = f"{chapter}. {title}"
        elif level >= 2 and chapter and not _NUMBERED.match(title):
            sub += 1
            title = f"{chapter}.{sub} {title}"
        body = assemble_section(
            original, normalize_typography(section["content"]), level
        )
        head, _, rest = body.partition("\n")
        lines = []
        for line in rest.split("\n"):
            inner = re.match(r"^##\s+(.*)$", line)
            if inner and chapter and level == 1:
                # Sub-headings the writer put inside a chapter join the same
                # numbering as the planned sub-sections that follow.
                sub += 1
                heading = re.sub(r"^\d+(?:\.\d+)*[.)]?\s+", "", inner.group(1))
                line = f"## {chapter}.{sub} {heading}"
            lines.append(line)
        parts.append(
            f"{'#' * level} {title}" + ("\n" + "\n".join(lines) if rest else "")
        )
    prefix = (language or "").lower()[:2]
    for kind, heading in (
        ("academic", bibliography_heading(language)),
        ("legal", LEGAL_HEADINGS.get(prefix, LEGAL_HEADINGS["en"])),
    ):
        rows = sorted(
            (r for r in bibliography if r.get("kind", "academic") == kind),
            key=lambda r: r.get("sort_key") or r["formatted"].casefold(),
        )
        if rows:
            parts.append(
                assemble_section(heading, "\n\n".join(r["formatted"] for r in rows))
            )
    return "\n\n".join(parts)


def add_table_of_contents(docx: Any, language: str | None) -> None:
    """A TOC field (levels 1-3) that Word fills in when the file is opened."""
    label = docx.add_paragraph()
    run = label.add_run(TOC_HEADINGS.get((language or "").lower()[:2], "Contents"))
    run.bold, run.font.size = True, Pt(14)
    paragraph = docx.add_paragraph()
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, 'TOC \\o "1-3" \\h \\z \\u'),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "…"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        run = paragraph.add_run()
        element = OxmlElement(tag)
        for name, value in attrs.items():
            element.set(qn(name), value)
        if text is not None:
            element.text = text
        run._r.append(element)
    settings = docx.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
    docx.add_page_break()


def apply_academic_profile(docx: Any) -> None:
    """Explicit neutral profile; never change type or spacing to hit page count."""
    for section in docx.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "List Bullet",
        "List Number",
    ):
        style = docx.styles[name]
        style.font.name = "Times New Roman"
        fonts = style.element.rPr.rFonts
        for attr in list(fonts.attrib):
            if "theme" in attr.casefold():
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), "Times New Roman")
        fonts.set(qn("w:cs"), "Times New Roman")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(
            16 if name == "Title" else 14 if name == "Heading 1" else 12
        )
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    docx.styles["Normal"].paragraph_format.widow_control = True
    docx.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def append_markdown(docx: Any, text: str) -> None:
    """Render CommonMark headings, emphasis and lists without losing visible text."""
    paragraph = None
    lists: list[str] = []
    bold = italic = 0
    link_href = ""
    link_label = ""
    for token in _PARSER.parse(text):
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            lists.append(
                "List Bullet" if token.type == "bullet_list_open" else "List Number"
            )
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "heading_open":
            paragraph = docx.add_paragraph(
                style=f"Heading {min(int(token.tag[1:]), 3)}"
            )
        elif token.type == "paragraph_open":
            paragraph = docx.add_paragraph(style=lists[-1] if lists else "Normal")
        elif token.type == "inline" and paragraph is not None:
            for child in token.children or []:
                if child.type == "strong_open":
                    bold += 1
                elif child.type == "strong_close":
                    bold -= 1
                elif child.type == "em_open":
                    italic += 1
                elif child.type == "em_close":
                    italic -= 1
                elif child.type in {"text", "code_inline", "html_inline", "image"}:
                    if link_href:
                        link_label += child.content
                    run = paragraph.add_run(child.content)
                    run.bold = True if bold else None
                    run.italic = True if italic else None
                elif child.type == "softbreak":
                    paragraph.add_run(" ")
                elif child.type == "hardbreak":
                    paragraph.add_run().add_break()
                elif child.type == "link_open":
                    link_href = str(child.attrGet("href") or "")
                    link_label = ""
                elif child.type == "link_close":
                    if link_href and link_label != link_href:
                        paragraph.add_run(f" ({link_href})")
                    link_href = ""
        elif token.type in {"fence", "code_block"}:
            docx.add_paragraph(token.content.rstrip())


def canonical_docx_bytes(data: bytes) -> bytes:
    """Remove ZIP container clock noise; the document's own metadata is retained."""
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

    output = BytesIO()
    with (
        ZipFile(BytesIO(data)) as archive,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for name in sorted(archive.namelist()):
            info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, archive.read(name))
    return output.getvalue()
