"""Regression evidence for the visible defects in work 7; no provider calls."""

import io
from contextlib import ExitStack

import pytest
from docx import Document as DocxDocument

from app.services.document_service import DocumentService
from tests.test_bibliography_export import _capture_upload, _seed_completed_document


@pytest.mark.asyncio
async def test_docx_converts_headings_emphasis_lists_and_preserves_references(
    db_session,
):
    content = (
        "# Introduzione\n\n## Metodo\nIl **confronto** è *critico*.\n\n"
        "### Criteri\n\n- Primo criterio\n- Secondo criterio\n\n"
        "1. Primo passo\n2. Secondo passo\n\n"
        "# Bibliografia\n\nRossi &amp; Bianchi (2024). Studio. "
        "https://doi.org/10.1234/a_b\n"
    )
    user, document = await _seed_completed_document(db_session, content=content)
    with ExitStack() as stack:
        uploads = _capture_upload(stack)
        await DocumentService(db_session).export_document(document.id, "docx", user.id)
    docx = DocxDocument(io.BytesIO(uploads["data"]))
    paragraphs = [(p.style.name, p.text) for p in docx.paragraphs]
    assert ("Heading 2", "Metodo") in paragraphs
    assert ("Heading 3", "Criteri") in paragraphs
    assert ("List Bullet", "Primo criterio") in paragraphs
    assert ("List Number", "Secondo passo") in paragraphs
    body = next(p for p in docx.paragraphs if "confronto" in p.text)
    assert body.text == "Il confronto è critico."
    assert any(r.bold and r.text == "confronto" for r in body.runs)
    assert any(r.italic and r.text == "critico" for r in body.runs)
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "Rossi & Bianchi (2024)" in text
    assert "https://doi.org/10.1234/a_b" in text
    assert "Sitografia" not in text


@pytest.mark.asyncio
async def test_sections_export_removes_only_leading_duplicate_title(db_session):
    content = "# Introduzione\n\nPrimo paragrafo.\n\n## Tema diverso\n\nAltro testo."
    user, document = await _seed_completed_document(
        db_session, sections=[("Introduzione", content, [])]
    )
    with ExitStack() as stack:
        uploads = _capture_upload(stack)
        await DocumentService(db_session).export_document(document.id, "docx", user.id)
    docx = DocxDocument(io.BytesIO(uploads["data"]))
    assert [p.text for p in docx.paragraphs].count("Introduzione") == 1
    assert any(
        p.text == "Tema diverso" and p.style.name == "Heading 2"
        for p in docx.paragraphs
    )
    assert "# Introduzione" not in "\n".join(p.text for p in docx.paragraphs)


@pytest.mark.asyncio
async def test_no_empty_sitography_and_explicit_default_page_profile(db_session):
    user, document = await _seed_completed_document(
        db_session, content="# Introduzione\n\nTesto."
    )
    with ExitStack() as stack:
        uploads = _capture_upload(stack)
        await DocumentService(db_session).export_document(document.id, "docx", user.id)
    docx = DocxDocument(io.BytesIO(uploads["data"]))
    assert "Sitografia" not in [p.text for p in docx.paragraphs]
    section = docx.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 1) == 2.5
    assert docx.styles["Normal"].font.name == "Times New Roman"
    assert docx.styles["Normal"].font.size.pt == 12
    assert docx.styles["Normal"].paragraph_format.line_spacing == 1.5
    assert not docx.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert "asciiTheme" not in docx.styles["Heading 1"].element.xml
    heading = next(p for p in docx.paragraphs if p.style.name == "Heading 1")
    assert all(run.bold is not False for run in heading.runs)


def test_assembly_is_idempotent_and_preserves_interior_titles_and_numbering():
    from app.services.docx_export import assemble_section

    body = "# Capitolo 1. Tema\n\nProsa.\n\n## 1.1 Dettaglio\n\n# Capitolo 1. Tema\n\nRipresa."
    assembled = assemble_section("Capitolo 1. Tema", body)
    assert assemble_section("Capitolo 1. Tema", assembled) == assembled
    assert assembled.count("# Capitolo 1. Tema") == 2
    assert "## 1.1 Dettaglio" in assembled


def test_markdown_links_keep_destination_and_literal_underscores():
    from app.services.docx_export import append_markdown

    doc = DocxDocument()
    append_markdown(
        doc, "[Studio](https://example.org/a_b) e https://doi.org/10.1234/a_b"
    )
    assert (
        doc.paragraphs[0].text
        == "Studio (https://example.org/a_b) e https://doi.org/10.1234/a_b"
    )
