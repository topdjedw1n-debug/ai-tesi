"""Italian traditional citations as Word footnotes (order 24286128, Vademecum)."""

import io
import re
import zipfile

from docx import Document as DocxDocument

from app.services.ai_pipeline.rag_retriever import SourceDoc
from app.services.citation_notes import NOTE_MARK, render_notes, surname
from app.services.docx_export import add_footnotes, append_markdown


def source(title, authors, year, **extra):
    return SourceDoc(
        title=title,
        authors=authors,
        year=year,
        canonical_metadata={"verification_provider": "PDF", "origin": "pdf"},
        **extra,
    )


KNOWN = {
    "KBERNARDI": source(
        "Il paesaggio nel cinema italiano", ["Sandro Bernardi"], 2002, venue="Marsilio"
    ),
    "KCHATMAN": source(
        "Antonioni: or, the Surface of the World", ["Seymour Chatman"], 1985
    ),
    "KBRUNETTA1": source("Storia del cinema italiano", ["Gian Piero Brunetta"], 1993),
    "KBRUNETTA2": source(
        "Il cinema neorealista italiano", ["Gian Piero Brunetta"], 2009
    ),
    "KWEB": source("Il Neorealismo '45-'51", ["Michele Corsi"], 2026),
    "KMANY": source(
        "Space &amp; image", ["A One", "B Two", "C Three", "D Four"], 2021, venue="J"
    ),
}
WEB = {
    "KWEB": {
        "site": "Cinescuola",
        "url": "https://www.cinescuola.it/x/",
        "accessed": "24 settembre 2026",
    }
}


def plain(text, notes):
    return NOTE_MARK.sub(lambda m: f"^{m.group(1)}", text)


def test_first_note_is_full_then_ibidem_ibid_and_op_cit():
    texts, notes, entries, missing = render_notes(
        [
            "Il vuoto è una forma [KBERNARDI] p. 12. Lo spazio pesa [KBERNARDI] p. 12. "
            "E si svuota [KBERNARDI] pp. 14-15.",
            "Il colore agisce [KCHATMAN] p. 3. Il paesaggio resta [KBERNARDI] p. 20.",
        ],
        KNOWN,
    )
    assert plain(texts[0], notes) == (
        "Il vuoto è una forma^1. Lo spazio pesa^2. E si svuota^3."
    )
    assert plain(texts[1], notes) == "Il colore agisce^4. Il paesaggio resta^5."
    assert notes == [
        "Sandro Bernardi, *Il paesaggio nel cinema italiano*, in «Marsilio», 2002, p. 12.",
        "Ibidem.",
        "Ibid., pp. 14–15.",
        "Seymour Chatman, *Antonioni: or, the Surface of the World*, 1985, p. 3.",
        "Bernardi, op. cit., p. 20.",
    ]
    assert set(entries) == {"KBERNARDI", "KCHATMAN"} and missing == []
    assert entries["KBERNARDI"]["formatted"].startswith(
        "Sandro Bernardi, *Il paesaggio"
    )


def test_same_surname_uses_short_title_cit():
    _, notes, _, _ = render_notes(
        [
            "A [KBRUNETTA1] p. 1. B [KBRUNETTA2] p. 2. C [KCHATMAN]. D [KBRUNETTA1] p. 5."
        ],
        KNOWN,
    )
    assert notes[3] == "Brunetta, *Storia del cinema italiano*, cit., p. 5."


def test_marker_as_noun_keeps_surnames_and_parentheses_fold():
    texts, notes, _, _ = render_notes(
        [
            "Secondo [KCHATMAN] p. 7, il deserto è interiore (cfr. testo) "
            "([KBERNARDI] p. 2; [KMANY] p. 4). [KBERNARDI] mostra il vuoto, "
            "affidato a [KCHATMAN]."
        ],
        KNOWN,
    )
    assert plain(texts[0], notes) == (
        "Secondo Chatman^1, il deserto è interiore (cfr. testo)^2. "
        "Bernardi^3 mostra il vuoto, affidato a Chatman^4."
    )
    assert notes[1] == (
        "Sandro Bernardi, *Il paesaggio nel cinema italiano*, in «Marsilio», 2002, p. 2; "
        "A One et al., *Space & image*, in «J», 2021, p. 4."
    )


def test_web_pages_are_cited_by_url_without_pages_and_unknown_keys_drop():
    texts, notes, entries, missing = render_notes(
        ["Il neorealismo nasce nel 1945 [KWEB] p. 1. Poi [KWEB] p. 2 e [KNOPE]."],
        KNOWN,
        WEB,
    )
    assert notes == [
        "Michele Corsi, *Il Neorealismo '45-'51*, in «Cinescuola», "
        "https://www.cinescuola.it/x/ (ultima consultazione: 24 settembre 2026).",
        "Ibidem.",
    ]
    assert plain(texts[0], notes) == "Il neorealismo nasce nel 1945^1. Poi^2 e."
    assert entries["KWEB"]["kind"] == "web" and missing == ["KNOPE"]


def test_key_written_as_a_word_becomes_the_authors_with_a_note():
    texts, notes, _, _ = render_notes(
        ["Un limite: le fonti KWEB e KCHATMAN sono divulgative [KWEB] p. 2."],
        KNOWN,
        WEB,
    )
    assert plain(texts[0], notes) == (
        "Un limite: le fonti Corsi^1 e Chatman^2 sono divulgative^3."
    )
    assert notes[2] == "Corsi, op. cit."


def test_repository_is_not_a_journal_and_subtitle_keeps_its_stop():
    thesis = source(
        "Redirecting Neorealism",
        ["Mary Lorraine DiSalvo"],
        2014,
        venue="Digital Access to Scholarship at Harvard (DASH) (Harvard University)",
        url="https://dash.harvard.edu/handle/1/12274297",
    )
    book = source(
        "L’estetica del vuoto\n Vuoti, pieni, suoni", ["Silvia Rivadossi"], 2023
    )
    _, notes, entries, _ = render_notes(
        ["A [KT] p. 3. B [KB]."], {"KT": thesis, "KB": book}
    )
    assert notes[0] == "Mary Lorraine DiSalvo, *Redirecting Neorealism*, 2014, p. 3."
    assert entries["KT"]["formatted"].endswith(
        "2014, https://dash.harvard.edu/handle/1/12274297."
    )
    assert notes[1].startswith("Silvia Rivadossi, *L’estetica del vuoto. Vuoti, pieni")
    kolker = source(
        "The Altering Eye", ["Robert Kolker"], 2009, venue="Open Book Publishers"
    )
    _, notes, _, _ = render_notes(["A [KK] p. 65."], {"KK": kolker})
    assert notes == [
        "Robert Kolker, *The Altering Eye*, Open Book Publishers, 2009, p. 65."
    ]
    record = source(
        "Tesi", ["Giulia Baso"], 2014, url="https://openalex.org/W2286797220"
    )
    _, _, entries, _ = render_notes(["A [KR]."], {"KR": record})
    assert entries["KR"]["formatted"] == "Giulia Baso, *Tesi*, 2014."


def test_surname_handles_particles_and_inverted_names():
    assert surname("Lorenzo De Luca") == "De Luca"
    assert surname("Brunetta, Gian Piero") == "Brunetta"


def test_footnotes_are_real_word_footnotes():
    texts, notes, _, _ = render_notes(["Il deserto *rosso* [KCHATMAN] p. 3."], KNOWN)
    docx = DocxDocument()
    append_markdown(docx, texts[0])
    add_footnotes(docx, notes)
    stream = io.BytesIO()
    docx.save(stream)
    archive = zipfile.ZipFile(io.BytesIO(stream.getvalue()))
    body = archive.read("word/document.xml").decode()
    notes_xml = archive.read("word/footnotes.xml").decode()
    assert '<w:footnoteReference w:id="1"/>' in body and "\ue000" not in body
    assert "footnotes+xml" in archive.read("[Content_Types].xml").decode()
    assert re.search(r'<w:footnote w:id="1">.*<w:i/>.*Antonioni', notes_xml)
    assert "<w:footnotePr>" in archive.read("word/settings.xml").decode()
    reopened = DocxDocument(io.BytesIO(stream.getvalue()))
    assert reopened.paragraphs[0].text == "Il deserto rosso."
