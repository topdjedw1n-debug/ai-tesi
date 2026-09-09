"""Partial real history stays partial; no synthetic response fills a missing record."""

import copy
import json
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

from app.services.academic_context import digest
from app.services.ai_pipeline.citation_formatter import CitationStyle
from app.services.brief_source_scopes import brief_source_scopes
from app.services.generation_policy import optional_stage_failure, warning_mode
from app.services.model_recording import ReplayIncomplete, ReplayTape
from app.services.outline_validation import validate_outline
from app.services.standard_references import render_standard_references

CASES = json.loads(
    (
        Path(__file__).parent / "fixtures/platform_first/historical_cases.json"
    ).read_text()
)


@pytest.mark.parametrize("case", CASES, ids=lambda case: f"historical-{case['work']}")
def test_available_history_preserves_structure_and_reports_incomplete_replay(case):
    before = digest(case)
    original = case["input"]["outline"]
    normalized = validate_outline(copy.deepcopy(original))
    assert [s["title"] for s in normalized["sections"]] == [
        s["title"] for s in original["sections"]
    ]
    # A real persisted outline is not its original raw provider response.
    assert case["origin"] == "historical_persisted_output_not_raw_model_response"
    job_ids = {
        e["payload"]["job_id"]
        for e in case["events"]
        if (e.get("payload") or {}).get("job_id")
    }
    tape = ReplayTape.from_events(case["events"], job_id=next(iter(job_ids), -1))
    with pytest.raises(ReplayIncomplete):
        tape.response(
            provider="anthropic", model="claude-opus-4-8", stage="outline", request={}
        )
    assert digest(case) == before


@pytest.mark.asyncio
async def test_real_brief_11_all_fourteen_nested_scopes_survive_search_expansion():
    case = next(c for c in CASES if c["work"] == 11)
    scopes = brief_source_scopes(
        case["input"]["additional_requirements"], case["input"]["outline"]
    )
    assert len(scopes) == 14
    assert any("termoregolazione" in scope for scope in scopes)
    assert any("NANDA-I" in scope for scope in scopes)
    assert "allattamento" in scopes[-1]
    from app.services.ai_pipeline.rag_retriever import RAGRetriever
    from app.services.ai_pipeline.source_pack import SourcePackBuilder

    rag = RAGRetriever()
    rag.search_crossref = AsyncMock(return_value=[])
    rag.search_openalex = AsyncMock(return_value=[])
    token = warning_mode.set(True)
    try:
        await SourcePackBuilder(rag).build(
            topic=case["input"]["topic"],
            language="it",
            document_id=11,
            section_titles=scopes,
        )
        queries = [call.args[0] for call in rag.search_crossref.call_args_list]
        assert len(set(queries)) >= len(scopes)
        assert any("termoregolazione" in query for query in queries)
        assert any("nanda" in query for query in queries)
        assert any(
            "allattamento" in query or "counseling" in query for query in queries
        )
        assert not optional_stage_failure(RuntimeError("unknown programming defect"))
    finally:
        warning_mode.reset(token)


def test_standard_reference_is_explicitly_unverified_and_never_keeps_invented_doi():
    raw = 'Definizione [STD:who].\n<STANDARD_REFERENCES_JSON>[{"id":"who","title":"Standard manual","authors":["WHO"],"year":null,"doi":"10.9999/invented","verified":true}]</STANDARD_REFERENCES_JSON>'
    prose, references, metadata = render_standard_references(raw, CitationStyle.APA)
    assert "STD:" not in prose and "STANDARD_REFERENCES_JSON" not in prose
    assert "WHO" in prose and "Standard manual" in references[0]
    assert (
        metadata[0]["identity_verified"] is False
        and metadata[0]["in_frozen_pack"] is False
    )
    assert "doi" not in metadata[0] and "10.9999" not in " ".join(references)


def test_standard_metadata_parse_failure_is_a_warning_without_metadata_in_prose():
    prose, references, metadata = render_standard_references(
        "Testo.\n<STANDARD_REFERENCES_JSON>{broken", CitationStyle.APA
    )
    assert prose == "Testo." and not references
    assert metadata[0]["status"] == "manager_verification_required"


def test_real_work_7_saved_text_exports_without_rewriting(tmp_path):
    from docx import Document as WordDocument

    from app.services.docx_export import (
        append_markdown,
        apply_academic_profile,
        assemble_section,
    )

    case = next(c for c in CASES if c["work"] == 7)
    before = digest(case)
    word = WordDocument()
    apply_academic_profile(word)
    for section in case["sections"]:
        append_markdown(word, assemble_section(section["title"], section["content"]))
    path = tmp_path / "historical-7-text-export.docx"
    word.save(path)
    reopened = WordDocument(path)
    headings = [p.text for p in reopened.paragraphs if p.style.name == "Heading 1"]
    assert all(section["title"] in headings for section in case["sections"])
    assert len(" ".join(p.text for p in reopened.paragraphs).split()) > 3000
    assert digest(case) == before


@pytest.mark.parametrize(
    "text",
    [
        "Testo [STD:nanda-i].",
        "Testo [STD:WHO2020].<STANDARD_REFERENCES_JSON>{broken",
        'Testo [STD:a].<STANDARD_REFERENCES_JSON>```json\n[{"id":"a","title":"Manuale","authors":["WHO"],"year":2020}]\n```</STANDARD_REFERENCES_JSON>',
    ],
)
def test_standard_transport_markers_never_reach_docx_prose(text):
    prose, bibliography, metadata = render_standard_references(text, CitationStyle.APA)
    assert "[STD:" not in prose and "STANDARD_REFERENCES_JSON" not in prose
    assert metadata


def test_grounded_outline_warning_prompt_preserves_uncovered_brief():
    from app.models.document import Document
    from app.services.ai_pipeline.source_pack import SourcePack
    from app.services.ai_service import AIService

    case = next(c for c in CASES if c["work"] == 11)
    document = Document(id=11, user_id=1, **case["input"])
    token = warning_mode.set(True)
    try:
        prompt = AIService(None)._build_grounded_outline_prompt(
            document, document.additional_requirements, SourcePack(11, document.topic)
        )
    finally:
        warning_mode.reset(token)
    assert "do not plan sections" not in prompt
    assert "preserve all required brief subsections" in prompt


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filename", ["real_writer_stage_11.json", "real_writer_stage_11_v2.json"]
)
async def test_recorded_real_writer_stage_11_replays_exact_prompt_without_sdk(
    db_session, monkeypatch, filename
):
    from datetime import datetime
    from types import SimpleNamespace
    from unittest.mock import MagicMock

    from sqlalchemy import DateTime

    from app.core.config import settings
    from app.models.auth import User
    from app.models.document import Document, DocumentSource
    from app.services.ai_pipeline.generator import SectionGenerator
    from app.services.model_recording import operation_section, replay_models
    from app.services.source_verification_stage import load_source_pack

    data = json.loads(
        (Path(__file__).parent / "fixtures/platform_first" / filename).read_text()
    )
    db_session.add(User(id=1, email="real-record-replay@example.invalid"))
    await db_session.commit()
    document = Document(
        id=11,
        user_id=1,
        ai_provider="anthropic",
        ai_model="claude-opus-4-8",
        **data["input"],
    )
    db_session.add(document)
    await db_session.commit()
    for row in data["source_rows"]:
        values = {
            c.name: (
                datetime.fromisoformat(row[c.name])
                if isinstance(c.type, DateTime) and isinstance(row[c.name], str)
                else row[c.name]
            )
            for c in DocumentSource.__table__.columns
            if c.name in row
        }
        db_session.add(DocumentSource(**values))
    await db_session.commit()
    pack = await load_source_pack(db_session, 11)
    sdk = SimpleNamespace(
        messages=SimpleNamespace(
            create=AsyncMock(side_effect=AssertionError("No live SDK call in replay"))
        ),
        close=AsyncMock(),
    )
    monkeypatch.setattr("anthropic.AsyncAnthropic", MagicMock(return_value=sdk))
    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "offline-placeholder")
    monkeypatch.setattr(settings, "AI_ENABLE_FALLBACK", False)
    monkeypatch.setattr(settings, "TRAINING_DATA_COLLECTION_ENABLED", False)
    tape = ReplayTape.from_events(data["document_provenance"], job_id=1)
    tape.allow_request_changes = filename == "real_writer_stage_11.json"
    # The first response tests rendering only; the final recording requires exact input.
    section = document.outline["sections"][data["section_index"] - 1]
    policy = warning_mode.set(True)
    section_token = operation_section.set(data["section_index"])
    try:
        with replay_models(tape):
            result = await SectionGenerator().generate_section(
                document,
                section["title"],
                data["section_index"],
                "anthropic",
                "claude-opus-4-8",
                citation_style=CitationStyle.APA,
                source_pack=pack,
                target_word_count=section["estimated_words"],
                additional_requirements=document.additional_requirements,
            )
            tape.assert_complete()
    finally:
        operation_section.reset(section_token)
        warning_mode.reset(policy)
    assert result["content"] == data["expected"]["content"].replace(
        "[Selix2015]", "(Selix, 2015)"
    )
    assert "[Selix2015]" not in result["content"]
    assert "Selix2015" in result["discarded_outline_keys"]
    assert tape.consumed[0]["request_changed"] is tape.allow_request_changes
    assert result["bibliography"] == data["expected"]["bibliography"]
    assert result["standard_references"] == data["expected"]["standard_references"]
    # A format correction is not source verification or a model rewrite.
    assert (
        result["unresolved_pack_markers"] == data["expected"]["unresolved_pack_markers"]
    )
    assert all(not r["identity_verified"] for r in result["standard_references"])
    sdk.messages.create.assert_not_called()


def test_outline_source_filter_preserves_every_brief_field_and_original():
    from app.services.ai_pipeline.prompt_builder import PromptBuilder

    case = next(c for c in CASES if c["work"] == 11)
    outline = copy.deepcopy(case["input"]["outline"])
    before = digest(outline)
    filtered, dropped = PromptBuilder.outline_for_available_sources(
        outline, {"Michel2017"}
    )
    assert "Selix2015" in dropped and digest(outline) == before

    def without_links(value):
        if isinstance(value, dict):
            return {
                k: without_links(v) for k, v in value.items() if k != "evidence_keys"
            }
        if isinstance(value, list):
            return [without_links(v) for v in value]
        return value

    assert without_links(filtered) == without_links(outline)
